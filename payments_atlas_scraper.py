#!/usr/bin/env python3
"""
Extrai os sistemas de pagamento de cada país do The Payments Atlas
(https://thepaymentsatlas.com) e guarda tudo num ficheiro .xlsx.

O site carrega os dados com JavaScript, por isso o script usa o Playwright
(um Chromium controlado por código) para abrir cada página como um browser real.

INSTALAÇÃO (uma vez):
    pip install playwright beautifulsoup4 openpyxl
    playwright install chromium

USO:
    python payments_atlas_para_xlsx.py                    # todos os países
    python payments_atlas_para_xlsx.py --codes AO,BO,CU   # só alguns países
    python payments_atlas_para_xlsx.py --no-expand        # sem clicar em "Show Full Details"
    python payments_atlas_para_xlsx.py --debug            # guarda o HTML da 1.ª página
    python payments_atlas_para_xlsx.py --from-html pagina.html --codes AO   # testar com HTML guardado

O progresso é guardado em payments_atlas_cache.json: se parar a meio, basta
voltar a correr o comando e ele continua de onde ficou (use --refresh para refazer).

Nota: use um intervalo razoável entre pedidos (--delay) e confirme os termos de
uso do site antes de extrair em massa.
"""
import argparse
import json
import os
import re
import sys
import time
from datetime import date
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = "https://thepaymentsatlas.com"
CACHE_FILE = "payments_atlas_cache.json"

# Linhas de interface que não fazem parte do conteúdo
RUIDO = {"show full details", "hide details", "show less", "show more", "official site"}
REF = re.compile(r"\[\d+\](?:\[\d+\])*")  # notas de rodapé tipo [107][115]


# --------------------------------------------------------------------------
# Parsing (funciona sobre HTML já renderizado)
# --------------------------------------------------------------------------
def limpar(s):
    s = s.replace("(opens in a new tab)", "")
    return re.sub(r"\s+", " ", s).strip()


def parse_item(li):
    site = ""
    for a in li.find_all("a", href=True):
        if "official site" in a.get_text(" ", strip=True).lower():
            site = urljoin(BASE, a["href"])
            break
    linhas = [limpar(x) for x in li.get_text("\n", strip=True).split("\n")]
    linhas = [x for x in linhas if x and x.lower() not in RUIDO]
    nome = linhas[0] if linhas else ""
    descricao = REF.sub("", linhas[1]).strip() if len(linhas) > 1 else ""
    detalhes = "\n".join(REF.sub("", x).strip() for x in linhas[2:])
    return {"nome": nome, "descricao": descricao, "site": site, "detalhes": detalhes}


def encontrar_lista(soup):
    """Devolve a <ul>/<ol> que vem logo a seguir ao título 'Payment Systems'."""
    for no in soup.find_all(string=re.compile(r"^\s*Payment Systems\s*$", re.I)):
        lista = no.find_next(["ul", "ol"])
        if lista and lista.find("li"):
            return lista
    return None


def parse_pais(html, codigo):
    soup = BeautifulSoup(html, "html.parser")
    linhas = [limpar(x) for x in soup.get_text("\n", strip=True).split("\n")]
    linhas = [x for x in linhas if x]

    nome, regiao, resumo = "", "", ""
    for i, l in enumerate(linhas):
        m = re.match(r"^(.+?)\s*·\s*([A-Z]{2})$", l)
        if m and m.group(2) == codigo.upper():
            regiao = m.group(1)
            nome = linhas[i + 1] if i + 1 < len(linhas) else ""
            resumo = linhas[i + 2] if i + 2 < len(linhas) else ""
            break
    if not nome:
        h1 = soup.find("h1")
        if h1:
            nome = re.sub(r"\s*Payment Systems\s*$", "", h1.get_text(" ", strip=True))

    declarado = None
    m = re.search(r"(\d+)\s+recorded", " ".join(linhas))
    if m:
        declarado = int(m.group(1))

    lista = encontrar_lista(soup)
    itens = lista.find_all("li", recursive=False) if lista else []
    sistemas = [s for s in (parse_item(li) for li in itens) if s["nome"]]

    return {
        "codigo": codigo.upper(),
        "nome": nome,
        "regiao": regiao,
        "resumo": resumo,
        "declarado": declarado,
        "url": f"{BASE}/country/{codigo.upper()}",
        "sistemas": sistemas,
    }


# --------------------------------------------------------------------------
# Obtenção das páginas (Playwright)
# --------------------------------------------------------------------------
ISO_2 = ("AD AE AF AG AI AL AM AO AQ AR AS AT AU AW AX AZ BA BB BD BE BF BG BH BI BJ BL BM BN BO BQ BR BS BT BV BW BY BZ "
         "CA CC CD CF CG CH CI CK CL CM CN CO CR CU CV CW CX CY CZ DE DJ DK DM DO DZ EC EE EG EH ER ES ET FI FJ FK FM FO FR "
         "GA GB GD GE GF GG GH GI GL GM GN GP GQ GR GS GT GU GW GY HK HM HN HR HT HU ID IE IL IM IN IO IQ IR IS IT JE JM JO JP "
         "KE KG KH KI KM KN KP KR KW KY KZ LA LB LC LI LK LR LS LT LU LV LY MA MC MD ME MF MG MH MK ML MM MN MO MP MQ MR MS MT "
         "MU MV MW MX MY MZ NA NC NE NF NG NI NL NO NP NR NU NZ OM PA PE PF PG PH PK PL PM PN PR PS PT PW PY QA RE RO RS RU RW "
         "SA SB SC SD SE SG SH SI SJ SK SL SM SN SO SR SS ST SV SX SY SZ TC TD TF TG TH TJ TK TL TM TN TO TR TT TV TW TZ UA UG "
         "UM US UY UZ VA VC VE VG VI VN VU WF WS XK YE YT ZA ZM ZW").split()
RE_PAIS = re.compile(r"/country/([A-Za-z]{2})(?:[/?#]|$)")


def _recolher_links(page, codigos):
    try:
        hrefs = page.eval_on_selector_all("a[href]", "els => els.map(e => e.href)")
    except Exception:
        hrefs = []
    for h in hrefs:
        m = RE_PAIS.search(h)
        if m:
            codigos.add(m.group(1).upper())


def obter_codigos(page):
    """Descobre os códigos dos países. Devolve (codigos, usou_lista_interna)."""
    codigos = set()
    try:
        page.goto(f"{BASE}/countries", wait_until="domcontentloaded", timeout=45000)
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:
        pass
    # fechar avisos de cookies, se existirem
    for txt in ("Accept all", "Accept", "I agree", "Got it", "OK"):
        try:
            page.get_by_role("button", name=txt).first.click(timeout=800)
            break
        except Exception:
            pass
    # descer a página para carregar listas preguiçosas, recolhendo links pelo caminho
    for _ in range(25):
        _recolher_links(page, codigos)
        page.mouse.wheel(0, 4000)
        page.wait_for_timeout(400)
    _recolher_links(page, codigos)

    # fallback 1: sitemap
    if not codigos:
        for caminho in ("/sitemap.xml", "/sitemap-0.xml"):
            try:
                page.goto(BASE + caminho, wait_until="domcontentloaded", timeout=20000)
                codigos.update(c.upper() for c in RE_PAIS.findall(page.content()))
            except Exception:
                pass
            if codigos:
                break

    if codigos:
        return sorted(codigos), False

    # fallback 2: lista interna de códigos ISO (as páginas inexistentes são ignoradas)
    try:
        page.goto(f"{BASE}/countries", wait_until="domcontentloaded", timeout=30000)
        with open("debug_countries.html", "w", encoding="utf-8") as f:
            f.write(page.content())
        print("Aviso: não encontrei links de países. Guardei debug_countries.html.")
    except Exception:
        pass
    print("A usar a lista interna de códigos ISO (mais lento; páginas inexistentes são ignoradas).")
    return sorted(ISO_2), True


def expandir_detalhes(page):
    """Clica em todos os 'Show Full Details'. Devolve {nome_do_sistema: texto_de_modal}."""
    extras = {}
    botoes = page.get_by_text("Show Full Details", exact=True).all()
    handles = []
    for b in botoes:
        try:
            handles.append(b.element_handle(timeout=1000))
        except Exception:
            pass
    for h in handles:
        try:
            nome = h.evaluate(
                "e => { const li = e.closest('li'); return li ? li.innerText.split('\\n')[0] : ''; }"
            )
            h.click(timeout=2000)
            page.wait_for_timeout(150)
            dialogo = page.locator("[role=dialog]")
            if dialogo.count():
                extras[nome.strip()] = limpar(dialogo.first.inner_text())
                page.keyboard.press("Escape")
                page.wait_for_timeout(100)
        except Exception:
            continue
    return extras


def obter_pais(page, codigo, expandir, tentativas=3, rapido=False):
    url = f"{BASE}/country/{codigo}"
    ultimo_erro = None
    if rapido:
        tentativas = 1
    for t in range(tentativas):
        try:
            resp = page.goto(url, wait_until="domcontentloaded", timeout=45000)
            if resp is not None and resp.status == 404:
                raise RuntimeError("página inexistente (404)")
            page.wait_for_selector("text=Payment Systems", timeout=8000 if rapido else 20000)
            page.wait_for_load_state("networkidle", timeout=10000 if rapido else 20000)
            extras = expandir_detalhes(page) if expandir else {}
            html = page.content()
            dados = parse_pais(html, codigo)
            for s in dados["sistemas"]:
                if s["nome"] in extras:
                    s["detalhes"] = (s["detalhes"] + "\n" + extras[s["nome"]]).strip()
            return dados, html
        except Exception as e:  # noqa: BLE001
            ultimo_erro = e
            time.sleep(2 * (t + 1))
    raise RuntimeError(f"{codigo}: {ultimo_erro}")


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
FONTE = "Arial"
COR_CAB = PatternFill("solid", start_color="1F3864")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def estilo_cabecalho(ws, n_cols):
    for c in range(1, n_cols + 1):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        cel.fill = COR_CAB
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDA
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ws.max_row}"


def formatar_corpo(ws, larguras, wrap_cols):
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.font = Font(name=FONTE, size=10)
            cel.border = BORDA
            cel.alignment = Alignment(vertical="top", wrap_text=cel.column in wrap_cols)


def gerar_xlsx(paises, caminho):
    wb = Workbook()

    # --- Folha 1: Sistemas (uma linha por sistema de pagamento) ---
    ws = wb.active
    ws.title = "Sistemas"
    cab = ["País", "Código", "Região", "Sistema de pagamento", "Descrição",
           "Site oficial", "Detalhes completos", "Página no Atlas"]
    ws.append(cab)
    for p in sorted(paises, key=lambda x: x["nome"] or x["codigo"]):
        for s in p["sistemas"]:
            ws.append([p["nome"], p["codigo"], p["regiao"], s["nome"], s["descricao"],
                       s["site"], s["detalhes"], p["url"]])
    for row in ws.iter_rows(min_row=2):
        for idx in (5, 7):  # colunas F e H -> hiperligações
            cel = row[idx]
            if cel.value:
                cel.hyperlink = cel.value
    estilo_cabecalho(ws, len(cab))
    formatar_corpo(ws, [22, 8, 14, 38, 50, 34, 60, 34], wrap_cols={4, 5, 7})
    for row in ws.iter_rows(min_row=2):
        for idx in (5, 7):
            row[idx].font = Font(name=FONTE, size=10, color="0563C1", underline="single")

    # --- Folha 2: Países (resumo) ---
    wp = wb.create_sheet("Países")
    cab2 = ["País", "Código", "Região", "Sistemas registados (segundo o site)",
            "Sistemas extraídos", "Resumo do site", "Página no Atlas"]
    wp.append(cab2)
    for i, p in enumerate(sorted(paises, key=lambda x: x["nome"] or x["codigo"]), start=2):
        wp.append([p["nome"], p["codigo"], p["regiao"], p["declarado"],
                   f'=COUNTIF(Sistemas!$B:$B,B{i})', p["resumo"], p["url"]])
        if p["url"]:
            wp.cell(row=i, column=7).hyperlink = p["url"]
    estilo_cabecalho(wp, len(cab2))
    formatar_corpo(wp, [22, 8, 14, 20, 14, 60, 34], wrap_cols={6})
    for row in wp.iter_rows(min_row=2):
        row[6].font = Font(name=FONTE, size=10, color="0563C1", underline="single")
        row[3].alignment = row[4].alignment = Alignment(horizontal="center", vertical="top")
        row[1].alignment = Alignment(horizontal="center", vertical="top")

    # --- Folha 3: Notas ---
    wn = wb.create_sheet("Notas")
    notas = [
        ("Fonte", "The Payments Atlas — https://thepaymentsatlas.com"),
        ("Data da extração", date.today().isoformat()),
        ("Conteúdo", "Folha 'Sistemas' = uma linha por sistema listado em 'Payment Systems' de cada país."),
        ("Conferência", "Na folha 'Países', compare 'registados' com 'extraídos'. Diferenças indicam itens por rever."),
        ("Referências [n]", "As notas de rodapé do tipo [107] foram removidas das descrições."),
        ("Direitos", "Os dados pertencem ao The Payments Atlas; confirme os termos de uso antes de redistribuir."),
    ]
    for k, v in notas:
        wn.append([k, v])
    for row in wn.iter_rows():
        row[0].font = Font(name=FONTE, bold=True, size=10)
        row[1].font = Font(name=FONTE, size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    wn.column_dimensions["A"].width = 20
    wn.column_dimensions["B"].width = 100

    wb.save(caminho)


# --------------------------------------------------------------------------
# Programa principal
# --------------------------------------------------------------------------
def carregar_cache():
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, encoding="utf-8") as f:
            return json.load(f)
    return {}


def guardar_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)


def main():
    ap = argparse.ArgumentParser(description="Payments Atlas -> xlsx")
    ap.add_argument("--codes", help="códigos ISO separados por vírgula (ex.: AO,BO). Omitir = todos")
    ap.add_argument("--out", default="payments_atlas.xlsx", help="ficheiro de saída")
    ap.add_argument("--delay", type=float, default=1.0, help="segundos entre países")
    ap.add_argument("--no-expand", action="store_true", help="não clicar em 'Show Full Details'")
    ap.add_argument("--refresh", action="store_true", help="ignorar a cache e refazer tudo")
    ap.add_argument("--debug", action="store_true", help="guardar o HTML do primeiro país")
    ap.add_argument("--show", action="store_true", help="mostrar o browser")
    ap.add_argument("--from-html", help="ler um HTML local em vez de abrir o site (usa --codes com 1 código)")
    args = ap.parse_args()

    # Modo offline: testar o parsing com um HTML guardado
    if args.from_html:
        if not args.codes:
            sys.exit("Indique o código do país com --codes (ex.: --codes AO)")
        with open(args.from_html, encoding="utf-8") as f:
            dados = parse_pais(f.read(), args.codes.split(",")[0])
        gerar_xlsx([dados], args.out)
        print(f"{dados['nome']}: {len(dados['sistemas'])} sistemas (site indica {dados['declarado']}) -> {args.out}")
        return

    from playwright.sync_api import sync_playwright

    cache = {} if args.refresh else carregar_cache()
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=not args.show)
        page = browser.new_page(viewport={"width": 1366, "height": 900})
        rapido = False
        if args.codes:
            codigos = [c.strip().upper() for c in args.codes.split(",") if c.strip()]
        else:
            codigos, rapido = obter_codigos(page)
            print(f"{len(codigos)} códigos a processar.")
        if not codigos:
            sys.exit("Nenhum código de país para processar.")

        falhas = []
        for n, c in enumerate(codigos, start=1):
            if c in cache:
                print(f"[{n}/{len(codigos)}] {c} (cache)")
                continue
            try:
                dados, html = obter_pais(page, c, expandir=not args.no_expand, rapido=rapido)
                cache[c] = dados
                guardar_cache(cache)
                if args.debug and n == 1:
                    with open(f"debug_{c}.html", "w", encoding="utf-8") as f:
                        f.write(html)
                aviso = ""
                if dados["declarado"] is not None and dados["declarado"] != len(dados["sistemas"]):
                    aviso = f"  ⚠ site indica {dados['declarado']}"
                print(f"[{n}/{len(codigos)}] {c} {dados['nome']}: {len(dados['sistemas'])} sistemas{aviso}")
            except Exception as e:  # noqa: BLE001
                if rapido:
                    print(f"[{n}/{len(codigos)}] {c} sem página")
                else:
                    falhas.append(c)
                    print(f"[{n}/{len(codigos)}] {c} FALHOU: {e}")
            time.sleep(args.delay)
        browser.close()

    paises = [cache[c] for c in codigos if c in cache]
    if not paises:
        sys.exit("Nenhum país extraído; xlsx não criado. Corra com --show --debug e envie o debug_*.html.")
    gerar_xlsx(paises, args.out)
    print(f"\nConcluído: {len(paises)} países, "
          f"{sum(len(p['sistemas']) for p in paises)} sistemas -> {args.out}")
    if falhas:
        print("Falharam (volte a correr o comando para tentar de novo):", ", ".join(falhas))


if __name__ == "__main__":
    main()