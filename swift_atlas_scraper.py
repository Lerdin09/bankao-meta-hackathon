#!/usr/bin/env python3
"""
Extrai os SWIFT BICs (aba "SWIFT BICs") de cada país do The Payments Atlas
e guarda tudo num xlsx.

REQUISITO: colocar este ficheiro na MESMA pasta que payments_atlas_scraper.py
(reutiliza a descoberta de países e as funções de formatação).

USO:
    python swift_atlas_scraper.py                    # todos os países
    python swift_atlas_scraper.py --codes US,AO      # só alguns
    python swift_atlas_scraper.py --show --debug     # ver o browser + guardar HTML
    python swift_atlas_scraper.py --from-html pagina.html --codes AO   # teste offline

Notas:
- Cada país mostra os BICs aos poucos ("Show More (N remaining)"); o script
  clica até aparecer tudo. Países grandes (EUA, etc.) demoram vários minutos.
- O progresso fica em cache_swift/<PAÍS>.json; se parar, volte a correr o
  comando e ele continua (use --refresh para refazer).
- Países com 0 BICs NÃO ficam em cache (voltam a ser tentados) e geram
  debug_swift_<PAÍS>.html para diagnóstico.
- Fonte indicada pelo site: ISO 9362 BIC records publicados pela OpenSanctions.
"""
import argparse
import json
import os
import re
import sys
import time
from datetime import date

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from payments_atlas_scraper import BASE, limpar, obter_codigos, parse_pais
except ImportError:
    sys.exit("Coloque este ficheiro na mesma pasta que payments_atlas_scraper.py")

PASTA_CACHE = "cache_swift"
FONTE = "Arial"
SEC_SEL = 'section[aria-labelledby="swift-heading"]'


# --------------------------------------------------------------------------
# Parsing (lê o DOM da secção SWIFT)
# --------------------------------------------------------------------------
def parse_bics(html):
    """Devolve (lista_de_registos, total_declarado)."""
    soup = BeautifulSoup(html, "html.parser")
    sec = soup.select_one(SEC_SEL)
    if sec is None:
        return [], None

    declarado = None
    for s in sec.find_all("span"):
        m = re.match(r"^([\d.,\s]+)\s*BICs?$", s.get_text(strip=True))
        if m:
            declarado = int(re.sub(r"\D", "", m.group(1)))
            break

    registos = []
    for li in sec.select("li"):
        code = li.find("code")
        if code is None:
            continue
        partes = [limpar(s.get_text(strip=True)) for s in code.find_all("span")]
        if len(partes) != 3:
            continue
        ps = li.find_all("p")
        lei = ""
        for s in li.find_all("span"):
            t = s.get_text(strip=True)
            if t.startswith("LEI "):
                lei = t[4:].strip()
        registos.append({
            "instituicao": limpar(ps[0].get_text(" ", strip=True)) if ps else "",
            "endereco": limpar(ps[1].get_text(" ", strip=True)) if len(ps) > 1 else "",
            "cod_inst": partes[0],
            "cod_pais": partes[1],
            "cod_loc": partes[2],
            "bic": "".join(partes),
            "lei": lei,
        })
    return registos, declarado


def extrair_regiao(html):
    """A região está no cabeçalho: <p class="eyebrow">Africa · AO</p>."""
    soup = BeautifulSoup(html, "html.parser")
    p = soup.select_one("p.eyebrow")
    if p is None:
        return ""
    return limpar(p.get_text(" ", strip=True).split("·")[0])


# --------------------------------------------------------------------------
# Playwright
# --------------------------------------------------------------------------
def abrir_aba_swift(page):
    tab = page.locator('[role="tab"][id$="-trigger-swift"]')
    if not tab.count():
        tab = page.get_by_role("tab", name=re.compile(r"SWIFT", re.I))
    if not tab.count():
        return False
    tab.first.click(timeout=5000)
    page.wait_for_selector("#swift-heading", timeout=15000)
    try:
        page.wait_for_selector(SEC_SEL + " li code", timeout=8000)
    except Exception:
        pass  # país com secção vazia
    return True


def carregar_tudo(page, max_cliques=5000):
    """Clica em 'Show More' (dentro da secção SWIFT) até não haver mais BICs."""
    sec = page.locator(SEC_SEL)
    anterior, paradas = None, 0
    for _ in range(max_cliques):
        btn = sec.get_by_role("button", name=re.compile(r"Show More"))
        if not btn.count():
            break
        try:
            txt = btn.first.inner_text(timeout=2000)
        except Exception:
            break
        m = re.search(r"\((\d[\d.,]*)\s+remaining\)", txt)
        restante = int(re.sub(r"\D", "", m.group(1))) if m else None
        if restante is not None and restante == anterior:
            paradas += 1
            if paradas >= 3:
                break
        else:
            paradas = 0
        anterior = restante
        try:
            btn.first.click(timeout=3000)
        except Exception:
            break
        page.wait_for_timeout(250)


def obter_bics(page, codigo, debug=False, rapido=False):
    url = f"{BASE}/country/{codigo}"
    resp = page.goto(url, wait_until="domcontentloaded", timeout=45000)
    if resp is not None and resp.status == 404:
        raise RuntimeError("página inexistente (404)")
    page.wait_for_selector("h1", timeout=8000 if rapido else 20000)
    page.wait_for_load_state("networkidle", timeout=10000 if rapido else 20000)
    html0 = page.content()
    cab = parse_pais(html0, codigo)  # só para o nome
    cab["regiao"] = extrair_regiao(html0) or cab.get("regiao") or ""

    tem_aba = abrir_aba_swift(page)
    if not tem_aba:
        return {"codigo": codigo, "nome": cab["nome"], "regiao": cab["regiao"],
                "declarado": 0, "url": url, "bics": [], "sem_aba": True}
    page.wait_for_timeout(500)
    carregar_tudo(page)
    html = page.content()
    bics, declarado = parse_bics(html)
    if debug or not bics:
        with open(f"debug_swift_{codigo}.html", "w", encoding="utf-8") as f:
            f.write(html)
    return {"codigo": codigo, "nome": cab["nome"], "regiao": cab["regiao"],
            "declarado": declarado, "url": url, "bics": bics, "sem_aba": False}


# --------------------------------------------------------------------------
# Excel (modo write-only: aguenta centenas de milhares de linhas)
# --------------------------------------------------------------------------
def _cel(ws, valor, negrito=False, cor=None, fundo=None, centro=False, wrap=False):
    c = WriteOnlyCell(ws, value=valor)
    c.font = Font(name=FONTE, size=10, bold=negrito, color=cor)
    if fundo:
        c.fill = PatternFill("solid", start_color=fundo)
    c.alignment = Alignment(horizontal="center" if centro else None,
                            vertical="top" if not centro else "center", wrap_text=wrap)
    return c


def _cabecalho(ws, nomes):
    ws.append([_cel(ws, n, negrito=True, cor="FFFFFF", fundo="1F3864", centro=True, wrap=True) for n in nomes])


def gerar_xlsx(paises, caminho):
    wb = Workbook(write_only=True)

    # Folha Países (resumo + conferência)
    wp = wb.create_sheet("Países")
    for col, w in zip("ABCDEF", [22, 8, 14, 18, 16, 40]):
        wp.column_dimensions[col].width = w
    wp.freeze_panes = "A2"
    _cabecalho(wp, ["País", "Código", "Região", "BICs (segundo o site)", "BICs extraídos", "Página no Atlas"])
    ordenados = sorted(paises, key=lambda p: p["nome"] or p["codigo"])
    for i, p in enumerate(ordenados, start=2):
        wp.append([
            _cel(wp, p["nome"]), _cel(wp, p["codigo"], centro=True), _cel(wp, p["regiao"]),
            _cel(wp, p["declarado"], centro=True),
            _cel(wp, f"=COUNTIF('SWIFT BICs'!$B:$B,B{i})", centro=True),
            _cel(wp, p["url"], cor="0563C1"),
        ])

    # Folha SWIFT BICs (uma linha por BIC)
    ws = wb.create_sheet("SWIFT BICs")
    for col, w in zip("ABCDEFGHIJ", [22, 8, 44, 60, 12, 12, 10, 12, 24, 34]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    _cabecalho(ws, ["País", "Código país (página)", "Instituição", "Endereço", "BIC (8)",
                    "Cód. instituição", "Cód. país (BIC)", "Cód. localização", "LEI", "Página no Atlas"])
    total = 0
    for p in ordenados:
        for b in p["bics"]:
            ws.append([
                _cel(ws, p["nome"]), _cel(ws, p["codigo"], centro=True), _cel(ws, b["instituicao"]),
                _cel(ws, b["endereco"]), _cel(ws, b["bic"], centro=True), _cel(ws, b["cod_inst"], centro=True),
                _cel(ws, b["cod_pais"], centro=True), _cel(ws, b["cod_loc"], centro=True),
                _cel(ws, b["lei"]), _cel(ws, p["url"], cor="0563C1"),
            ])
            total += 1
    if total > 1_048_000:
        print("⚠ Mais linhas do que o Excel suporta numa folha; divida a extração por --codes.")

    # Folha Notas
    wn = wb.create_sheet("Notas")
    wn.column_dimensions["A"].width = 20
    wn.column_dimensions["B"].width = 100
    for k, v in [
        ("Fonte", "The Payments Atlas — https://thepaymentsatlas.com (aba SWIFT BICs de cada país)"),
        ("Origem dos dados", "Segundo o site: registos ISO 9362 BIC publicados pela OpenSanctions."),
        ("Data da extração", date.today().isoformat()),
        ("BIC (8)", "Junção de código da instituição + código do país + código de localização, como mostrado no site."),
        ("Conferência", "Na folha 'Países' compare 'segundo o site' com 'extraídos'; diferenças indicam registos em falta."),
        ("Direitos", "Os dados pertencem às fontes originais; confirme os termos de uso antes de redistribuir."),
    ]:
        wn.append([_cel(wn, k, negrito=True), _cel(wn, v, wrap=True)])

    wb.save(caminho)


# --------------------------------------------------------------------------
# Programa principal
# --------------------------------------------------------------------------
def cache_path(codigo):
    return os.path.join(PASTA_CACHE, f"{codigo}.json")


def main():
    ap = argparse.ArgumentParser(description="Payments Atlas SWIFT BICs -> xlsx")
    ap.add_argument("--codes", help="códigos ISO separados por vírgula. Omitir = todos")
    ap.add_argument("--out", default="payments_atlas_swift.xlsx")
    ap.add_argument("--delay", type=float, default=1.0)
    ap.add_argument("--refresh", action="store_true")
    ap.add_argument("--debug", action="store_true")
    ap.add_argument("--show", action="store_true")
    ap.add_argument("--from-html", help="HTML local da aba SWIFT BICs (teste offline; usa --codes com 1 código)")
    args = ap.parse_args()

    if args.from_html:
        if not args.codes:
            sys.exit("Indique o código do país com --codes (ex.: --codes AO)")
        with open(args.from_html, encoding="utf-8") as f:
            html = f.read()
        cab = parse_pais(html, args.codes.split(",")[0])
        cab["regiao"] = extrair_regiao(html) or cab.get("regiao") or ""
        bics, declarado = parse_bics(html)
        gerar_xlsx([{"codigo": cab["codigo"], "nome": cab["nome"], "regiao": cab["regiao"],
                     "declarado": declarado, "url": cab["url"], "bics": bics}], args.out)
        print(f"{cab['nome']}: {len(bics)} BICs (site indica {declarado}) -> {args.out}")
        return

    from playwright.sync_api import sync_playwright

    os.makedirs(PASTA_CACHE, exist_ok=True)
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=not args.show)
        page = browser.new_page(viewport={"width": 1366, "height": 900})
        rapido = False
        if args.codes:
            codigos = [c.strip().upper() for c in args.codes.split(",") if c.strip()]
        else:
            codigos, rapido = obter_codigos(page)
            print(f"{len(codigos)} códigos a processar.")
        if not codigos:
            sys.exit("Nenhum código de país para processar.")

        falhas = []
        try:
            for n, c in enumerate(codigos, start=1):
                if os.path.exists(cache_path(c)) and not args.refresh:
                    with open(cache_path(c), encoding="utf-8") as f:
                        d = json.load(f)
                    if not d.get("regiao"):
                        try:
                            page.goto(f"{BASE}/country/{c}", wait_until="domcontentloaded", timeout=45000)
                            page.wait_for_selector("p.eyebrow", timeout=15000)
                            d["regiao"] = extrair_regiao(page.content())
                            with open(cache_path(c), "w", encoding="utf-8") as f:
                                json.dump(d, f, ensure_ascii=False)
                            print(f"[{n}/{len(codigos)}] {c} (cache) região: {d['regiao']}")
                            time.sleep(args.delay)
                            continue
                        except Exception as e:  # noqa: BLE001
                            print(f"[{n}/{len(codigos)}] {c} (cache) região falhou: {e}")
                            continue
                    print(f"[{n}/{len(codigos)}] {c} (cache)")
                    continue
                try:
                    dados = obter_bics(page, c, debug=args.debug and n == 1, rapido=rapido)
                    # só guarda em cache se houver BICs (evita "envenenar" a cache com zeros)
                    if dados["bics"]:
                        with open(cache_path(c), "w", encoding="utf-8") as f:
                            json.dump(dados, f, ensure_ascii=False)
                    aviso = ""
                    if dados["sem_aba"]:
                        aviso = "  (sem aba SWIFT BICs)"
                    elif not dados["bics"]:
                        aviso = f"  (0 BICs; site indica {dados['declarado']}; ver debug_swift_{c}.html)"
                    elif dados["declarado"] is not None and dados["declarado"] != len(dados["bics"]):
                        aviso = f"  ⚠ site indica {dados['declarado']}"
                    print(f"[{n}/{len(codigos)}] {c} {dados['nome']}: {len(dados['bics'])} BICs{aviso}")
                except KeyboardInterrupt:
                    raise
                except Exception as e:  # noqa: BLE001
                    if rapido:
                        print(f"[{n}/{len(codigos)}] {c} sem página")
                    else:
                        falhas.append(c)
                        print(f"[{n}/{len(codigos)}] {c} FALHOU: {e}")
                time.sleep(args.delay)
        except KeyboardInterrupt:
            print("\nInterrompido. A gerar o xlsx com o que já existe em cache...")
        finally:
            try:
                browser.close()
            except Exception:  # noqa: BLE001
                pass

    paises = []
    for c in codigos:
        if os.path.exists(cache_path(c)):
            with open(cache_path(c), encoding="utf-8") as f:
                paises.append(json.load(f))
    paises = [p for p in paises if p["bics"]]
    if not paises:
        sys.exit("Nenhum BIC extraído; xlsx não criado. Veja os ficheiros debug_swift_*.html.")
    gerar_xlsx(paises, args.out)
    print(f"\nConcluído: {len(paises)} países, {sum(len(p['bics']) for p in paises)} BICs -> {args.out}")
    if falhas:
        print("Falharam (volte a correr o comando):", ", ".join(falhas))


if __name__ == "__main__":
    main()